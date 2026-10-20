"""
Extract plain text from an uploaded file so it can go through the same
chunking/indexing path regardless of what format the user attached.

Supports: .txt / .md (read as-is), .docx (python-docx), .pdf (pypdf).
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


def extract_text(filename: str, raw_bytes: bytes) -> str:
    """The file's text. Raises ValueError (shown to the user as a 400) for an
    unsupported type or a file that can't be read."""
    ext = Path(filename).suffix.lower()
    if ext in (".docx", ".pdf"):
        # A corrupt, encrypted or mislabelled file makes the parsers raise all
        # sorts of errors; report it as unreadable rather than a 500.
        try:
            return _extract_document(ext, raw_bytes)
        except Exception:
            raise ValueError(f"Couldn't read {filename}: the file looks damaged or isn't a valid "
                             f"{ext[1:].upper()} (password-protected files aren't supported).")

    if ext in (".txt", ".md"):
        return raw_bytes.decode("utf-8", errors="ignore")

    raise ValueError(
        f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _extract_document(ext: str, raw_bytes: bytes) -> str:
    if ext == ".docx":
        doc = DocxDocument(io.BytesIO(raw_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also pull text out of tables, since Word docs often put real content there.
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n\n".join(parts)

    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(raw_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(p for p in pages if p.strip())
